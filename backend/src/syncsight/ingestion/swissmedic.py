"""Swissmedic establishment licences (medicinal products incl. blood), one row per site.

Source: https://opendata.swiss/en/dataset/betriebsbewilligungen (OGD "open use").
Holders carry no UID; they are matched to Zefix companies by name, then postal code.
"""

import io
import logging
import re
import warnings
from collections.abc import Iterable, Sequence
from dataclasses import asdict, dataclass, field
from datetime import date, datetime
from typing import Any

import httpx2
from openpyxl import load_workbook
from sqlalchemy import delete, select, update
from sqlalchemy.dialects.postgresql import insert
from sqlalchemy.orm import Session

from syncsight.db.models import SwissmedicSite, SwissmedicSiteChange
from syncsight.ingestion import http
from syncsight.ingestion.matching import link_by_name
from syncsight.ingestion.normalize import name_key

log = logging.getLogger(__name__)

DOWNLOAD_URL = "https://ogd.swissmedic.cloud/swissmedic/OGD/bewilligung/SwissGMDP-Export_MP.xlsx"

# German part of each header cell (text before the first newline) -> record field
_COLUMNS = {
    "Bewilligungsnummer": "licence_number",
    "Inhaberin Betriebsbewilligung - (gem. Handelsregister)": "holder_name",
    "Strasse (Inhaberin)": "holder_street",
    "Hausnummer (Inhaberin)": "holder_house_number",
    "PLZ (Inhaberin)": "holder_postal_code",
    "Ort (Inhaberin)": "holder_locality",
    "Kanton (Inhaberin)": "holder_canton",
    "Gültig bis:": "valid_until",
    "Standort-Nummer": "site_number",
    "Name (Standort)": "site_name",
    "Strasse (Standort)": "site_street",
    "Hausnummer (Standort)": "site_house_number",
    "PLZ (Standort)": "site_postal_code",
    "Ort (Standort)": "site_locality",
    "Kanton (Standort)": "site_canton",
    "Zertifikat Nummer": "certificate_number",
    "Datum letzte Inspektion": "last_inspection",
    "Compliance-Status Standort": "compliance_status",
}
_ACTIVITY_HEADER = re.compile(r"^([1-8]) ")  # "1 Herstellung", ..., "8 Blut und labile ..."
# Licence activities (export column headers, translated)
ACTIVITIES = {
    1: "Manufacturing",
    2: "Import",
    3: "API manufacturing",
    4: "Wholesale",
    5: "Export",
    6: "Trade abroad",
    7: "Broker or agent",
    8: "Blood products",
}


@dataclass(frozen=True)
class SiteRecord:
    licence_number: str
    site_number: str
    holder_name: str
    holder_street: str | None
    holder_postal_code: str | None
    holder_locality: str | None
    holder_canton: str | None
    valid_until: str | None
    site_name: str | None
    site_street: str | None
    site_postal_code: str | None
    site_locality: str | None
    site_canton: str | None
    certificate_number: str | None
    last_inspection: date | None
    compliance_status: str | None
    activity_codes: list[int] = field(default_factory=list)
    activity_details: dict[str, str] = field(default_factory=dict)


def parse_workbook(content: bytes) -> list[SiteRecord]:
    with warnings.catch_warnings():  # the export has no default style; harmless
        warnings.simplefilter("ignore", UserWarning)
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    rows = list(workbook.worksheets[0].iter_rows(values_only=True))
    header_index = next(
        (i for i, row in enumerate(rows) if row and _header(row[0]) == "Bewilligungsnummer"), None
    )
    if header_index is None:
        raise ValueError("Swissmedic workbook: header row 'Bewilligungsnummer' not found")
    columns = [_header(h) for h in rows[header_index]]
    return [
        _record(dict(zip(columns, row, strict=False)))
        for row in rows[header_index + 1 :]
        if row and row[0]
    ]


def download(client: httpx2.Client) -> bytes:
    return http.request(client, "GET", DOWNLOAD_URL).content


def replace_sites(session: Session, records: Sequence[SiteRecord]) -> int:
    """The export is a full snapshot: record what changed since the last one, then replace it."""
    previous = {
        (s.licence_number, s.site_number): s for s in session.scalars(select(SwissmedicSite))
    }
    if previous:
        session.add_all(diff_sites(previous, records, detected_on=date.today()))
    session.execute(delete(SwissmedicSite))
    rows = [asdict(r) | {"holder_name_key": name_key(r.holder_name)} for r in records]
    if rows:
        session.execute(insert(SwissmedicSite).values(rows))
    return len(rows)


def diff_sites(
    previous: dict[tuple[str, str], SwissmedicSite],
    records: Sequence[SiteRecord],
    detected_on: date,
) -> list[SwissmedicSiteChange]:
    """Sites added or removed, activities added or removed, compliance status changed."""
    changes = []

    def change(record: SiteRecord | SwissmedicSite, kind: str, detail: str) -> None:
        changes.append(
            SwissmedicSiteChange(
                detected_on=detected_on,
                licence_number=record.licence_number,
                site_number=record.site_number,
                company_uid=previous[key].company_uid if key in previous else None,
                kind=kind,
                detail=detail,
            )
        )

    current = {(r.licence_number, r.site_number): r for r in records}
    for key, record in current.items():
        old = previous.get(key)
        place = f"{record.site_name}, {record.site_locality}"
        if old is None:
            change(record, "site_added", f"{place}: {_activity_names(record.activity_codes)}")
            continue
        if added := set(record.activity_codes) - set(old.activity_codes):
            change(record, "activity_added", f"{place}: {_activity_names(added)}")
        if removed := set(old.activity_codes) - set(record.activity_codes):
            change(record, "activity_removed", f"{place}: {_activity_names(removed)}")
        if record.compliance_status != old.compliance_status:
            detail = f"{place}: {old.compliance_status} → {record.compliance_status}"
            change(record, "compliance_changed", detail)
    for key, old in previous.items():
        if key not in current:
            change(old, "site_removed", f"{old.site_name}, {old.site_locality}")
    return changes


def match_companies(session: Session) -> int:
    """Link holders to the one active company with the same name (postal code breaks ties)."""
    holders = session.execute(
        select(SwissmedicSite.holder_name_key, SwissmedicSite.holder_postal_code).distinct()
    ).all()
    links = link_by_name(session, ((key, plz) for key, plz in holders))
    session.execute(update(SwissmedicSite).values(company_uid=None))
    unlinked_changes = SwissmedicSiteChange.company_uid.is_(None)
    for (key, postal_code), uid in links.items():
        session.execute(
            update(SwissmedicSite)
            .where(
                SwissmedicSite.holder_name_key == key,
                SwissmedicSite.holder_postal_code.is_not_distinct_from(postal_code),
            )
            .values(company_uid=uid)
        )
    # changes detected before the site was linked inherit the link
    session.execute(
        update(SwissmedicSiteChange)
        .where(
            unlinked_changes, SwissmedicSiteChange.licence_number == SwissmedicSite.licence_number
        )
        .values(company_uid=SwissmedicSite.company_uid)
    )
    return len(links)


def run(session: Session, client: httpx2.Client) -> tuple[int, int]:
    records = parse_workbook(download(client))
    stored = replace_sites(session, records)
    matched = match_companies(session)
    session.commit()
    log.info("swissmedic: %d sites, %d holders matched to companies", stored, matched)
    return stored, matched


def _record(row: dict[str | None, Any]) -> SiteRecord:
    v = {field: _value(row.get(header)) for header, field in _COLUMNS.items()}
    activities = {
        match.group(1): text
        for header, cell in row.items()
        if (match := _ACTIVITY_HEADER.match(header or "")) and (text := _value(cell))
    }
    inspection = v["last_inspection"]
    return SiteRecord(
        licence_number=v["licence_number"] or "",
        site_number=v["site_number"] or "",
        holder_name=v["holder_name"] or "",
        holder_street=_join(v["holder_street"], v["holder_house_number"]),
        holder_postal_code=v["holder_postal_code"],
        holder_locality=v["holder_locality"],
        holder_canton=v["holder_canton"],
        valid_until=v["valid_until"],
        site_name=v["site_name"],
        site_street=_join(v["site_street"], v["site_house_number"]),
        site_postal_code=v["site_postal_code"],
        site_locality=v["site_locality"],
        site_canton=v["site_canton"],
        certificate_number=v["certificate_number"],
        last_inspection=datetime.strptime(inspection, "%d.%m.%Y").date() if inspection else None,
        compliance_status=v["compliance_status"],
        activity_codes=sorted(int(code) for code in activities),
        activity_details=activities,
    )


def _header(value: Any) -> str | None:
    """Header cells are bilingual ('Ort (Inhaberin)\n\nLieu (siège)'); keep the German part."""
    text = _value(value)
    return text.split("\n", 1)[0].strip() if text else None


def _value(value: Any) -> str | None:
    if isinstance(value, float) and value.is_integer():
        value = int(value)  # numbers such as site numbers arrive as floats
    text = str(value).strip() if value is not None else ""
    return text or None


def _join(street: str | None, number: str | None) -> str | None:
    return " ".join(part for part in (street, number) if part) or None


def matched_uids(session: Session) -> list[str]:
    return list(
        session.scalars(
            select(SwissmedicSite.company_uid)
            .where(SwissmedicSite.company_uid.is_not(None))
            .distinct()
            .order_by(SwissmedicSite.company_uid)
        )
    )


def _activity_names(codes: Iterable[int]) -> str:
    return ", ".join(ACTIVITIES.get(code, str(code)) for code in sorted(codes))
